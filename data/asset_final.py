"""
Resolución de «Asset Final» (avance Col. 4) → assets de estadísticas en BD.

Formato en Excel: COLOMBIA-{region}-{version} (hoja MAPA GENERAL COLOMBIA, col. CJ).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from config import (
    ASSET_PARENT,
    AVANCE_COLOMBIA_XLSX,
    AVANCE_COL_ASSET_FINAL,
    AVANCE_SHEET_MAPA_GENERAL,
)
from data.db import get_conn, ph

_RE_ASSET_FINAL = re.compile(r"^COLOMBIA-(\d+)-(\d+)$", re.IGNORECASE)


@dataclass
class ResolucionAssetFinal:
    """Resultado al cruzar Asset Final con la BD para un bioma."""

    asset_ids: list[str] = field(default_factory=list)
    faltantes: list[str] = field(default_factory=list)  # no hay stats en BD
    fuera_bioma: list[str] = field(default_factory=list)
    invalidos: list[str] = field(default_factory=list)
    fuente: str = ""


def _score_leaf(leaf: str) -> tuple[int, int]:
    """Prioriza MapaGeneral > filtro-espacial > base exacta > resto."""
    low = leaf.lower().replace("-", "_")
    if "mapageneral" in low or "mapa_general" in low:
        prio = 0
    elif "espacial" in low:
        prio = 1
    elif re.match(r"^r\d+_v\d+$", low):
        prio = 2
    else:
        prio = 3
    return (prio, len(leaf))


def _prefijo_stats() -> str:
    return ASSET_PARENT.rstrip("/") + "/"


@lru_cache(maxsize=4)
def leer_asset_final_desde_xlsx(
    ruta: str | None = None,
    hoja: str = AVANCE_SHEET_MAPA_GENERAL,
    columna: str = AVANCE_COL_ASSET_FINAL,
) -> tuple[str, ...]:
    """Lee valores no vacíos de la columna Asset Final (orden del Excel)."""
    path = Path(ruta) if ruta else AVANCE_COLOMBIA_XLSX
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo de avance: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    if hoja not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Hoja no encontrada: {hoja}. Hojas: {wb.sheetnames}")

    ws = wb[hoja]
    col_idx = column_index_from_string(columna)
    out: list[str] = []
    vistos: set[str] = set()
    for row in ws.iter_rows(min_row=1, min_col=col_idx, max_col=col_idx, values_only=True):
        raw = row[0]
        if raw is None:
            continue
        val = str(raw).strip()
        if not val or val.lower().startswith("asset final"):
            continue
        if val in vistos:
            continue
        vistos.add(val)
        out.append(val)
    wb.close()
    return tuple(out)


def parse_asset_final(label: str) -> tuple[str, int] | None:
    """COLOMBIA-30205-8 → ('30205', 8)."""
    m = _RE_ASSET_FINAL.match(str(label).strip())
    if not m:
        return None
    return m.group(1), int(m.group(2))


def _buscar_candidatos_stats(region_id: str, version: int) -> list[tuple[str, str]]:
    """[(asset_id, bioma)] que coinciden R{region}_V{version}; prioriza ASSET_PARENT."""
    conn = get_conn()
    cur = conn.cursor()
    cur.execute(
        f"""
        SELECT asset_id, bioma FROM assets
        WHERE region_id = {ph()}
        """,
        (str(region_id),),
    )
    rows = cur.fetchall()
    conn.close()

    pat = re.compile(
        rf"^R0*{re.escape(str(region_id))}[_-]V0*{version}(?:$|[_-].*)$",
        re.IGNORECASE,
    )
    prefijo = _prefijo_stats()
    hits: list[tuple[str, str]] = []
    for aid, bioma in rows:
        leaf = aid.rsplit("/", 1)[-1]
        if pat.match(leaf):
            hits.append((aid, bioma or ""))

    def _rank(item: tuple[str, str]) -> tuple[int, int, int]:
        aid, _ = item
        bajo_parent = 0 if aid.startswith(prefijo) else 1
        prio, length = _score_leaf(aid.rsplit("/", 1)[-1])
        return (bajo_parent, prio, length)

    hits.sort(key=_rank)
    return hits


def resolver_assets_bioma_desde_avance(
    bioma_nombre: str,
    ruta_xlsx: str | Path | None = None,
) -> ResolucionAssetFinal:
    """
    Cruza Asset Final del Excel con la BD para el bioma indicado.

    - Si el asset no está en BD → faltantes (aviso, no error).
    - Si está pero es de otro bioma → fuera_bioma.
    - Una región / una versión: si hay varias filas, se conserva la de mayor versión.
    """
    ruta = str(ruta_xlsx) if ruta_xlsx else None
    try:
        labels = leer_asset_final_desde_xlsx(ruta)
        fuente = str(Path(ruta) if ruta else AVANCE_COLOMBIA_XLSX)
    except FileNotFoundError as exc:
        return ResolucionAssetFinal(faltantes=[str(exc)], fuente=str(AVANCE_COLOMBIA_XLSX))
    except Exception as exc:  # noqa: BLE001 — UI amigable
        return ResolucionAssetFinal(faltantes=[f"Error leyendo avance: {exc}"], fuente=str(AVANCE_COLOMBIA_XLSX))

    # region_id → (version, asset_id, label)
    por_region: dict[str, tuple[int, str, str]] = {}
    res = ResolucionAssetFinal(fuente=fuente)

    for label in labels:
        parsed = parse_asset_final(label)
        if not parsed:
            res.invalidos.append(label)
            continue
        rid, ver = parsed
        candidatos = _buscar_candidatos_stats(rid, ver)
        if not candidatos:
            res.faltantes.append(label)
            continue

        # Preferir match del bioma pedido
        del_bioma = [(a, b) for a, b in candidatos if b == bioma_nombre]
        if not del_bioma:
            res.fuera_bioma.append(label)
            continue

        asset_id = del_bioma[0][0]
        prev = por_region.get(rid)
        if prev is None or ver > prev[0]:
            por_region[rid] = (ver, asset_id, label)

    res.asset_ids = [t[1] for t in sorted(por_region.values(), key=lambda t: t[2])]
    return res
