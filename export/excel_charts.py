"""
Exportación Excel con gráficas (flujo original Colab / openpyxl).

Una hoja por versión + gráfico general e individuales con colores MapBiomas.
Nombre de descarga esperado: region_{id}_complete.xlsx
"""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

COLUMNAS_EXCLUIR = ["system:index", "descripcion", "version", ".geo", "geo"]

# Colores MapBiomas (script original)
COLORES_ID = {
    1: "1F8D49",
    3: "1F8D49",
    5: "04381D",
    6: "026975",
    49: "02D659",
    10: "D6BC74",
    11: "519799",
    12: "D6BC74",
    32: "FC8114",
    29: "FFAA5F",
    50: "AD5100",
    14: "FFEFC3",
    9: "7A5900",
    35: "9065D0",
    74: "BE83F7",
    21: "FFEFC3",
    22: "D4271E",
    23: "FFA07A",
    24: "D4271E",
    30: "9C0027",
    68: "E97A7A",
    25: "DB4D4F",
    75: "C12100",
    26: "2532E4",
    33: "2532E4",
    31: "091077",
    34: "93DFE6",
    27: "000000",
}


def _id_desde_header(header) -> int | None:
    if header is None:
        return None
    s = str(header).strip()
    m = re.match(r"^ID0*(\d+)$", s, re.IGNORECASE)
    if m:
        return int(m.group(1))
    if s.isdigit():
        return int(s)
    return None


def _normalizar_cols_id(df: pd.DataFrame) -> pd.DataFrame:
    ren = {}
    for c in df.columns:
        if c in ("year", "version"):
            continue
        id_val = _id_desde_header(c)
        if id_val is not None:
            ren[c] = f"ID{id_val:02d}" if id_val < 100 else f"ID{id_val}"
    out = df.rename(columns=ren)
    return out.loc[:, ~out.columns.duplicated()]


def _nombre_proceso_hoja(asset_o_label: str) -> str:
    """Deriva nombre de hoja tipo GAPFILL_V2 / CLASIFICACION_ORIGINAL_V1."""
    label = str(asset_o_label).rsplit("/", 1)[-1].replace("-", "_")

    m_ready = re.match(r"^([A-Za-zÁÉÍÓÚÑ_]+)_V(\d+)$", label, re.IGNORECASE)
    if m_ready:
        return f"{m_ready.group(1).upper()}_V{m_ready.group(2)}"

    m = re.search(r"_V(\d+)[_-]?(.*)$", label, re.IGNORECASE)
    if not m:
        return _sheet_name(label, set())

    version = m.group(1)
    sufijo = (m.group(2) or "").lower().strip("_")

    if "join" in sufijo:
        nombre = "JOIN"
    elif "gapfill" in sufijo:
        nombre = "GAPFILL"
    elif "frecuencia" in sufijo:
        nombre = "FRECUENCIA"
    elif "temporal" in sufijo:
        nombre = "TEMPORAL"
    elif "espacial" in sufijo:
        nombre = "ESPACIAL"
    elif "mapageneral" in sufijo or "mapa_general" in sufijo:
        nombre = "MAPAGENERAL"
    elif "clasificacion" in sufijo or sufijo in ("", "v1"):
        nombre = "CLASIFICACION_ORIGINAL"
    elif sufijo:
        nombre = re.sub(r"[^A-Za-z0-9]+", "_", sufijo).strip("_").upper()
    else:
        nombre = "BASE"

    return f"{nombre}_V{version}"


def _sheet_name(nombre: str, usados: set[str]) -> str:
    limpio = re.sub(r"[\\/*?:\[\]]", "_", str(nombre))[:31] or "Hoja"
    base = limpio
    i = 1
    while limpio in usados:
        suf = f"_{i}"
        limpio = base[: 31 - len(suf)] + suf
        i += 1
    usados.add(limpio)
    return limpio


def _df_para_hoja(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = out.drop(columns=[c for c in COLUMNAS_EXCLUIR if c in out.columns], errors="ignore")
    out = _normalizar_cols_id(out)
    if "year" not in out.columns:
        raise ValueError("El DataFrame no tiene columna year")
    id_cols = sorted(
        [c for c in out.columns if _id_desde_header(c) is not None],
        key=lambda c: _id_desde_header(c) or 0,
    )
    out = out[["year"] + id_cols].sort_values("year")
    out = out.fillna(0)
    out["year"] = out["year"].astype(int)
    for c in id_cols:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0)
    return out


def _limpiar_formato_grafico(chart) -> None:
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.majorGridlines = None
    chart.x_axis.majorGridlines = None
    chart.x_axis.tickLblSkip = 0


def _agregar_graficos_hoja(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    if "year" not in headers:
        return

    col_year = headers.index("year") + 1
    columnas_id = [
        i + 1 for i, col in enumerate(headers) if str(col).startswith("ID")
    ]
    if not columnas_id:
        return

    max_row = ws.max_row
    cats = Reference(ws, min_col=col_year, min_row=2, max_row=max_row)

    # 1) Gráfico general
    chart_all = LineChart()
    chart_all.title = "Evolución de Coberturas"
    chart_all.height, chart_all.width = 8, 38
    if chart_all.legend is not None:
        chart_all.legend.position = "b"

    for col_idx in columnas_id:
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=max_row)
        chart_all.add_data(data, titles_from_data=True)

    chart_all.set_categories(cats)
    _limpiar_formato_grafico(chart_all)

    for i, col_idx in enumerate(columnas_id):
        header = ws.cell(row=1, column=col_idx).value
        id_val = _id_desde_header(header) or 0
        color = COLORES_ID.get(id_val, "000000")
        if i < len(chart_all.series):
            serie = chart_all.series[i]
            serie.graphicalProperties.line.solidFill = color
            serie.graphicalProperties.line.width = 20000

    ws.add_chart(chart_all, "H2")

    # 2) Individuales
    start_row = 20
    for idx, col_idx in enumerate(columnas_id):
        header = ws.cell(row=1, column=col_idx).value
        id_val = _id_desde_header(header) or 0
        color = COLORES_ID.get(id_val, "000000")

        c = LineChart()
        c.title = str(header)
        c.height, c.width = 6, 18
        c.legend = None

        d = Reference(ws, min_col=col_idx, min_row=1, max_row=max_row)
        c.add_data(d, titles_from_data=True)
        c.set_categories(cats)
        _limpiar_formato_grafico(c)

        if c.series:
            serie = c.series[0]
            serie.graphicalProperties.line.solidFill = color
            serie.graphicalProperties.line.width = 20000

        col_pos = "H" if idx % 2 == 0 else "Z"
        row_pos = start_row + (idx // 2) * 15
        ws.add_chart(c, f"{col_pos}{row_pos}")


def generar_excel_con_graficas_desde_data_dict(
    data_dict: Mapping[str, pd.DataFrame],
) -> bytes:
    """Genera un .xlsx en memoria: una hoja por entrada + gráficas (script original)."""
    if not data_dict:
        raise ValueError("No hay datos para exportar")

    usados: set[str] = set()
    limpios: dict[str, pd.DataFrame] = {}
    for nombre, df in data_dict.items():
        hoja = _sheet_name(_nombre_proceso_hoja(nombre), usados)
        limpios[hoja] = _df_para_hoja(df)

    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "complete.xlsx"
        with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
            for hoja, df in limpios.items():
                df.to_excel(writer, sheet_name=hoja, index=False)

        wb = load_workbook(ruta)
        for hoja in wb.sheetnames:
            _agregar_graficos_hoja(wb[hoja])
        wb.save(ruta)
        return ruta.read_bytes()


def generar_excel_con_graficas_desde_region_xlsx(ruta_o_bytes) -> bytes:
    """Limpia REGIÓN_*.xlsx y añade gráficas (utilidad / tests)."""
    xls = pd.ExcelFile(ruta_o_bytes)
    data = {}
    for hoja in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=hoja)
        df = df.drop(columns=[c for c in COLUMNAS_EXCLUIR if c in df.columns])
        data[hoja] = df
    return generar_excel_con_graficas_desde_data_dict(data)
