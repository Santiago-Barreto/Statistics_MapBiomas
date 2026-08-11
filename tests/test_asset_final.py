"""Tests de Asset Final (avance Col. 4) → estadísticas BD."""

from pathlib import Path

import sqlite3

import pytest
from openpyxl import Workbook

from data import asset_final as af_mod
from data.asset_final import (
    leer_asset_final_desde_xlsx,
    parse_asset_final,
    resolver_assets_bioma_desde_avance,
)
from tests.conftest import _bootstrap_schema


@pytest.fixture
def temp_db(monkeypatch, tmp_path: Path):
    db_path = str(tmp_path / "test_asset_final.db")
    monkeypatch.setattr("data.db.DB_PATH", db_path)
    conn = sqlite3.connect(db_path)
    _bootstrap_schema(conn)
    conn.close()
    return db_path


def test_parse_asset_final():
    assert parse_asset_final("COLOMBIA-30205-8") == ("30205", 8)
    assert parse_asset_final("colombia-30102-6") == ("30102", 6)
    assert parse_asset_final("R30205_V8") is None


def test_leer_asset_final_desde_xlsx(tmp_path):
    path = tmp_path / "avance.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "MAPA GENERAL COLOMBIA"
    ws["CJ2"] = "Asset final "
    ws["CJ5"] = "COLOMBIA-30205-8"
    ws["CJ6"] = "COLOMBIA-30208-8"
    ws["CJ7"] = "COLOMBIA-30205-8"  # dup
    wb.save(path)

    af_mod.leer_asset_final_desde_xlsx.cache_clear()
    out = leer_asset_final_desde_xlsx(str(path))
    assert out == ("COLOMBIA-30205-8", "COLOMBIA-30208-8")


def test_resolver_assets_bioma_desde_avance(temp_db, tmp_path):
    from data.db import get_conn

    conn = get_conn()
    rows = [
        (
            "projects/mapbiomas-colombia/assets/LULC/COLECCION4/ESTADISTICAS/R30205_V8-filtro-espacial",
            "30205",
            "Amazonia",
            "R30205_V8-filtro-espacial",
            1,
        ),
        (
            "projects/mapbiomas-colombia/assets/LULC/COLECCION4/ESTADISTICAS/R30102_V6-filtro-espacial",
            "30102",
            "Andes",
            "R30102_V6-filtro-espacial",
            1,
        ),
    ]
    conn.executemany(
        "INSERT INTO assets (asset_id, region_id, bioma, label, last_sync) VALUES (?,?,?,?,?)",
        rows,
    )
    conn.commit()
    conn.close()

    path = tmp_path / "avance.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "MAPA GENERAL COLOMBIA"
    ws["CJ2"] = "Asset final"
    ws["CJ5"] = "COLOMBIA-30205-8"
    ws["CJ6"] = "COLOMBIA-30999-1"  # faltante
    ws["CJ7"] = "COLOMBIA-30102-6"  # otro bioma
    ws["CJ8"] = "NO-VALIDO"
    wb.save(path)

    af_mod.leer_asset_final_desde_xlsx.cache_clear()
    res = resolver_assets_bioma_desde_avance("Amazonia", ruta_xlsx=path)
    assert len(res.asset_ids) == 1
    assert "R30205_V8-filtro-espacial" in res.asset_ids[0]
    assert "COLOMBIA-30999-1" in res.faltantes
    assert "COLOMBIA-30102-6" in res.fuera_bioma
    assert "NO-VALIDO" in res.invalidos


def test_preferir_version_mayor_misma_region(temp_db, tmp_path):
    from data.db import get_conn

    conn = get_conn()
    conn.executemany(
        "INSERT INTO assets (asset_id, region_id, bioma, label, last_sync) VALUES (?,?,?,?,?)",
        [
            (
                "projects/x/assets/ESTADISTICAS/R30205_V6-filtro-espacial",
                "30205",
                "Amazonia",
                "a",
                1,
            ),
            (
                "projects/x/assets/ESTADISTICAS/R30205_V8-filtro-espacial",
                "30205",
                "Amazonia",
                "b",
                1,
            ),
        ],
    )
    conn.commit()
    conn.close()

    path = tmp_path / "avance.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "MAPA GENERAL COLOMBIA"
    ws["CJ5"] = "COLOMBIA-30205-6"
    ws["CJ6"] = "COLOMBIA-30205-8"
    wb.save(path)

    af_mod.leer_asset_final_desde_xlsx.cache_clear()
    res = resolver_assets_bioma_desde_avance("Amazonia", ruta_xlsx=path)
    assert len(res.asset_ids) == 1
    assert "V8" in res.asset_ids[0]
