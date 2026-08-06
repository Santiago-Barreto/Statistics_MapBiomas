"""Pruebas de exportación Excel con gráficas."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from export.excel_charts import (
    generar_excel_con_graficas_desde_data_dict,
    generar_excel_con_graficas_desde_region_xlsx,
)


def test_generar_desde_data_dict_crea_hojas_y_charts():
    df = pd.DataFrame(
        {
            "year": [1986, 1987, 1988],
            "ID03": [100.0, 110.0, 120.0],
            "ID21": [200.0, 190.0, 180.0],
            "version": ["V1", "V1", "V1"],
        }
    )
    raw = generar_excel_con_graficas_desde_data_dict({"CLASIF_V1": df})
    assert isinstance(raw, (bytes, bytearray))
    assert len(raw) > 1000

    wb = load_workbook(BytesIO(raw))
    assert "CLASIF_V1" in wb.sheetnames
    ws = wb["CLASIF_V1"]
    assert ws.max_row >= 4
    assert len(ws._charts) >= 1


def test_generar_desde_region_xlsx(tmp_path):
    path = tmp_path / "REGIÓN_30450.xlsx"
    df = pd.DataFrame(
        {
            "system:index": [0, 1],
            "ID03": [1.0, 2.0],
            "ID21": [3.0, 4.0],
            "descripcion": ["a", "b"],
            "version": [1, 1],
            "year": [1986, 1987],
            ".geo": ["{}", "{}"],
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="GAPFILL_V2", index=False)

    raw = generar_excel_con_graficas_desde_region_xlsx(path)
    wb = load_workbook(BytesIO(raw))
    assert "GAPFILL_V2" in wb.sheetnames
    assert "system:index" not in [
        c.value for c in wb["GAPFILL_V2"][1]
    ]
    assert len(wb["GAPFILL_V2"]._charts) >= 1
